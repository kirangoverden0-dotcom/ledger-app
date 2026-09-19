import io
import os
from datetime import date, datetime, timedelta
from functools import wraps

from sqlalchemy import func, case

# pyrefly: ignore [missing-import]
from flask import Flask, jsonify, request, session, send_from_directory, send_file

from models import db, Retailer, Transaction, Setting, BOOKS, IST, now_ist

app = Flask(__name__, static_folder="static", static_url_path="")
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-me")

db_url = os.environ.get("DATABASE_URL")
if db_url and db_url.strip():
    db_url = db_url.strip()
    # Render/Heroku-style postgres URLs sometimes start with postgres:// which
    # SQLAlchemy 1.4+ no longer accepts — normalize it.
    if db_url.startswith("postgres://"):
        db_url = db_url.replace("postgres://", "postgresql://", 1)
else:
    db_url = "sqlite:///ledger.db"

app.config["SQLALCHEMY_DATABASE_URI"] = db_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
    "pool_pre_ping": True,
    "pool_recycle": 280,
}
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


@app.after_request
def add_header(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


db.init_app(app)
with app.app_context():
    db.create_all()


def get_active_pin():
    # 1. Environment variable if set
    env_pin = os.environ.get("LEDGER_PIN")
    if env_pin and env_pin.strip():
        return env_pin.strip()
    # 2. Database Setting if present
    setting = Setting.query.get("pin")
    if setting and setting.value:
        return setting.value
    # No fallback PIN!
    return None


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("authed"):
            return jsonify({"error": "unauthorized"}), 401
        return f(*args, **kwargs)

    return wrapper


def valid_book(book):
    return book in BOOKS


# ---------------- Auth & PIN Management ----------------


@app.route("/api/pin/status")
def pin_status():
    active_pin = get_active_pin()
    return jsonify({"is_set": bool(active_pin)})


@app.route("/api/pin/setup", methods=["POST"])
def pin_setup():
    active_pin = get_active_pin()
    if active_pin is not None:
        return jsonify({"error": "PIN already set"}), 400

    data = request.get_json(silent=True) or {}
    new_pin = str(data.get("pin", "")).strip()
    if not new_pin:
        return jsonify({"error": "PIN required"}), 400

    setting = Setting.query.get("pin")
    if not setting:
        setting = Setting(key="pin", value=new_pin)
        db.session.add(setting)
    else:
        setting.value = new_pin
    db.session.commit()

    session["authed"] = True
    session.permanent = True
    return jsonify({"ok": True})


@app.route("/api/login", methods=["POST"])
def login():
    active_pin = get_active_pin()
    if active_pin is None:
        return jsonify({"error": "pin_not_set", "is_set": False}), 400

    data = request.get_json(silent=True) or {}
    pin_input = str(data.get("pin", "")).strip()

    if pin_input == active_pin:
        session["authed"] = True
        session.permanent = True
        return jsonify({"ok": True})
    return jsonify({"error": "wrong pin"}), 401


@app.route("/api/pin/change", methods=["POST"])
@login_required
def pin_change():
    data = request.get_json(silent=True) or {}
    new_pin = str(data.get("new_pin", "")).strip()
    if not new_pin:
        return jsonify({"error": "New PIN required"}), 400

    setting = Setting.query.get("pin")
    if not setting:
        setting = Setting(key="pin", value=new_pin)
        db.session.add(setting)
    else:
        setting.value = new_pin
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/session")
def session_check():
    return jsonify({
        "authed": bool(session.get("authed")),
        "pin_set": bool(get_active_pin())
    })


# ---------------- Summary ----------------


@app.route("/api/summary")
@login_required
def summary():
    """Returns total balance and retailer count for each book in a single aggregated query per book."""
    out = {}
    for book in BOOKS:
        book_filter = Retailer.has_santhoor if book == "santhoor" else Retailer.has_mtr

        tx_balance = func.coalesce(
            func.sum(
                case(
                    (Transaction.type == "purchase", Transaction.amount),
                    (Transaction.type == "payment", -Transaction.amount),
                    else_=0.0,
                )
            ),
            0.0,
        )

        subq = (
            db.session.query(
                Retailer.id,
                tx_balance.label("retailer_balance"),
            )
            .outerjoin(
                Transaction,
                (Transaction.retailer_id == Retailer.id) & (Transaction.book == book),
            )
            .filter(book_filter == True)
            .group_by(Retailer.id)
            .subquery()
        )

        row = db.session.query(
            func.count(subq.c.id),
            func.coalesce(func.sum(subq.c.retailer_balance), 0.0),
        ).one()

        count, total = row
        out[book] = {"total": round(float(total or 0.0), 2), "count": count}
    return jsonify(out)


# ---------------- Retailers ----------------


@app.route("/api/retailers")
@login_required
def list_retailers():
    """Returns all retailers in requested book with calculated balances using a single aggregated SQL query."""
    book = request.args.get("book", "santhoor")
    if not valid_book(book):
        return jsonify({"error": "invalid book"}), 400

    book_filter = Retailer.has_santhoor if book == "santhoor" else Retailer.has_mtr

    tx_balance = func.coalesce(
        func.sum(
            case(
                (Transaction.type == "purchase", Transaction.amount),
                (Transaction.type == "payment", -Transaction.amount),
                else_=0.0,
            )
        ),
        0.0,
    )

    results = (
        db.session.query(
            Retailer.id,
            Retailer.name,
            Retailer.has_santhoor,
            Retailer.has_mtr,
            tx_balance.label("balance"),
        )
        .outerjoin(
            Transaction,
            (Transaction.retailer_id == Retailer.id) & (Transaction.book == book),
        )
        .filter(book_filter == True)
        .group_by(Retailer.id, Retailer.name, Retailer.has_santhoor, Retailer.has_mtr)
        .order_by(Retailer.name.asc())
        .all()
    )

    out = [
        {
            "id": r_id,
            "name": name,
            "has_santhoor": s_flag,
            "has_mtr": m_flag,
            "balance": round(float(bal or 0.0), 2),
        }
        for r_id, name, s_flag, m_flag, bal in results
    ]
    return jsonify(out)


@app.route("/api/retailers", methods=["POST"])
@login_required
def create_retailer():
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "name required"}), 400

    has_santhoor = bool(data.get("has_santhoor", True))
    has_mtr = bool(data.get("has_mtr", True))

    if not has_santhoor and not has_mtr:
        return jsonify({"error": "Retailer must belong to at least one product line"}), 400

    r = Retailer(name=name, has_santhoor=has_santhoor, has_mtr=has_mtr)
    db.session.add(r)
    db.session.flush()

    for book in BOOKS:
        if not r.is_in_book(book):
            continue
        opening = data.get(f"opening_{book}")
        if opening:
            try:
                amt = float(opening)
            except (TypeError, ValueError):
                amt = 0
            if amt > 0:
                db.session.add(
                    Transaction(
                        retailer_id=r.id,
                        book=book,
                        type="purchase",
                        amount=amt,
                        note="Opening balance",
                    )
                )
    db.session.commit()
    return jsonify({"id": r.id, "name": r.name, "has_santhoor": r.has_santhoor, "has_mtr": r.has_mtr})


@app.route("/api/retailers/<int:retailer_id>")
@login_required
def get_retailer(retailer_id):
    book = request.args.get("book", "santhoor")
    if not valid_book(book):
        return jsonify({"error": "invalid book"}), 400
    r = Retailer.query.get_or_404(retailer_id)
    return jsonify(
        {
            "id": r.id,
            "name": r.name,
            "has_santhoor": r.has_santhoor,
            "has_mtr": r.has_mtr,
            "balance": r.balance(book),
            "history": r.history(book),
        }
    )


@app.route("/api/retailers/<int:retailer_id>", methods=["PUT"])
@login_required
def update_retailer(retailer_id):
    r = Retailer.query.get_or_404(retailer_id)
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "name required"}), 400

    has_santhoor = bool(data.get("has_santhoor", r.has_santhoor))
    has_mtr = bool(data.get("has_mtr", r.has_mtr))

    if not has_santhoor and not has_mtr:
        return jsonify({"error": "Retailer must belong to at least one product line"}), 400

    r.name = name
    r.has_santhoor = has_santhoor
    r.has_mtr = has_mtr
    db.session.commit()
    return jsonify({"id": r.id, "name": r.name, "has_santhoor": r.has_santhoor, "has_mtr": r.has_mtr})


@app.route("/api/retailers/<int:retailer_id>", methods=["DELETE"])
@login_required
def delete_retailer(retailer_id):
    r = Retailer.query.get_or_404(retailer_id)
    tx_count = Transaction.query.filter_by(retailer_id=r.id).count()
    if tx_count > 0:
        return jsonify({
            "error": f"Cannot delete '{r.name}' because they have {tx_count} transaction record(s). Deletion blocked to preserve transaction history.",
            "has_history": True,
            "tx_count": tx_count
        }), 400

    db.session.delete(r)
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/api/retailers/bulk-import", methods=["POST"])
@login_required
def bulk_import_retailers():
    """Bulk import retailers from uploaded .xlsx or .json file or JSON payload.
    Intelligently maps column headers, skips duplicates, and returns import summary.
    """
    rows_to_process = []
    columns_found = []
    mapping_override = None

    if request.is_json:
        payload = request.get_json(silent=True) or {}
        rows_to_process = payload.get("rows", [])
        mapping_override = payload.get("mapping")
    elif "file" in request.files:
        file = request.files["file"]
        filename = file.filename.lower()

        if filename.endswith(".json"):
            try:
                content = json.load(file)
                if isinstance(content, list):
                    rows_to_process = content
                elif isinstance(content, dict):
                    rows_to_process = content.get("retailers") or content.get("rows") or content.get("data") or []
            except Exception as e:
                return jsonify({"error": f"Invalid JSON file: {str(e)}"}), 400

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file, data_only=True)
                sheet = wb.active
                data = list(sheet.iter_rows(values_only=True))
                if not data:
                    return jsonify({"error": "Excel file is empty"}), 400

                raw_headers = [str(cell or "").strip() for cell in data[0]]
                columns_found = [h for h in raw_headers if h]

                for row_idx in range(1, len(data)):
                    row = data[row_idx]
                    if not any(row):
                        continue
                    row_dict = {}
                    for col_idx, col_name in enumerate(raw_headers):
                        if col_name and col_idx < len(row):
                            row_dict[col_name] = row[col_idx]
                    rows_to_process.append(row_dict)
            except ImportError:
                return jsonify({"error": "openpyxl library is required to read .xlsx files. Run 'pip install openpyxl'"}), 500
            except Exception as e:
                return jsonify({"error": f"Error parsing Excel file: {str(e)}"}), 400
        else:
            return jsonify({"error": "Unsupported file format. Please upload an .xlsx or .json file."}), 400

    if not rows_to_process:
        return jsonify({"error": "No retailer records found in file or payload"}), 400

    def match_column(col_names, keywords):
        for col in col_names:
            col_lower = str(col).lower().replace("_", " ").replace("-", " ").strip()
            for kw in keywords:
                if kw in col_lower:
                    return col
        return None

    # Handle simple list of string names e.g. ["Shop 1", "Shop 2"]
    if isinstance(rows_to_process[0], str):
        rows_to_process = [{"name": s} for s in rows_to_process]

    field_map = {}
    if isinstance(rows_to_process[0], dict):
        sample_keys = list(rows_to_process[0].keys())
        if not columns_found:
            columns_found = sample_keys

        detected_name_col = match_column(sample_keys, ["retailer name", "shop name", "name", "retailer", "shop", "store"])
        detected_santhoor_col = match_column(sample_keys, ["santoor", "santhoor"])
        detected_mtr_col = match_column(sample_keys, ["mtr"])
        detected_open_santhoor_col = match_column(sample_keys, ["opening santhoor", "opening santoor", "santhoor balance", "santoor balance", "santhoor opening", "santoor opening"])
        detected_open_mtr_col = match_column(sample_keys, ["opening mtr", "mtr balance", "mtr opening"])

        if mapping_override:
            field_map = mapping_override
        else:
            confirm_required = request.args.get("confirm") == "1" or not detected_name_col
            if confirm_required and len(sample_keys) > 1 and not (detected_name_col and (detected_santhoor_col or detected_mtr_col)):
                return jsonify({
                    "status": "mapping_required",
                    "columns": sample_keys,
                    "suggested_mapping": {
                        "name": detected_name_col or sample_keys[0],
                        "has_santhoor": detected_santhoor_col or "",
                        "has_mtr": detected_mtr_col or "",
                        "opening_santhoor": detected_open_santhoor_col or "",
                        "opening_mtr": detected_open_mtr_col or ""
                    },
                    "preview_rows": rows_to_process[:3]
                })

            field_map = {
                "name": detected_name_col or sample_keys[0],
                "has_santhoor": detected_santhoor_col or "",
                "has_mtr": detected_mtr_col or "",
                "opening_santhoor": detected_open_santhoor_col or "",
                "opening_mtr": detected_open_mtr_col or ""
            }

    existing_rows = db.session.query(Retailer.name).all()
    existing_names_set = {r[0].strip().lower() for r in existing_rows}

    added_list = []
    skipped_list = []
    batch_pending = 0

    for row in rows_to_process:
        if isinstance(row, dict):
            name_key = field_map.get("name") or "name"
            raw_name = str(row.get(name_key, "") or "").strip()
        else:
            raw_name = str(row).strip()

        if not raw_name:
            continue

        raw_name_lower = raw_name.lower()
        if raw_name_lower in existing_names_set:
            skipped_list.append(raw_name)
            continue

        existing_names_set.add(raw_name_lower)

        has_santhoor = True
        has_mtr = True

        if isinstance(row, dict):
            s_col = field_map.get("has_santhoor")
            if s_col and s_col in row and row[s_col] is not None:
                val = str(row[s_col]).lower().strip()
                has_santhoor = val in ["true", "1", "yes", "y", "checked"]

            m_col = field_map.get("has_mtr")
            if m_col and m_col in row and row[m_col] is not None:
                val = str(row[m_col]).lower().strip()
                has_mtr = val in ["true", "1", "yes", "y", "checked"]

            # Line check e.g. "product_line" column
            line_val = str(row.get("product_line") or row.get("line") or row.get("products") or "").lower()
            if line_val:
                has_santhoor = "santoor" in line_val or "santhoor" in line_val or "both" in line_val or "all" in line_val
                has_mtr = "mtr" in line_val or "both" in line_val or "all" in line_val

        if not has_santhoor and not has_mtr:
            has_santhoor = True
            has_mtr = True

        r = Retailer(name=raw_name, has_santhoor=has_santhoor, has_mtr=has_mtr)
        db.session.add(r)
        db.session.flush()

        if isinstance(row, dict):
            for b_key in ["santhoor", "mtr"]:
                col_k = field_map.get(f"opening_{b_key}")
                if col_k and col_k in row and row[col_k] is not None:
                    try:
                        amt = float(row[col_k])
                    except (ValueError, TypeError):
                        amt = 0
                    if amt > 0 and r.is_in_book(b_key):
                        db.session.add(
                            Transaction(
                                retailer_id=r.id,
                                book=b_key,
                                type="purchase",
                                amount=amt,
                                note="Opening balance",
                            )
                        )

        added_list.append(raw_name)
        batch_pending += 1

        if batch_pending >= 50:
            db.session.commit()
            db.session.expunge_all()
            batch_pending = 0

    if batch_pending > 0:
        db.session.commit()
        db.session.expunge_all()


    return jsonify({
        "status": "success",
        "added_count": len(added_list),
        "skipped_count": len(skipped_list),
        "added": added_list,
        "skipped": skipped_list
    })



# ---------------- Transactions ----------------


@app.route("/api/transactions", methods=["POST"])
@login_required
def create_transaction():
    data = request.get_json(silent=True) or {}
    retailer_id = data.get("retailer_id")
    book = data.get("book")
    ttype = data.get("type")
    amount = data.get("amount")
    entry_date = data.get("date")  # YYYY-MM-DD optional custom timestamp

    if not valid_book(book):
        return jsonify({"error": "invalid book"}), 400
    if ttype not in ("purchase", "payment"):
        return jsonify({"error": "invalid type"}), 400
    try:
        amount = float(amount)
    except (TypeError, ValueError):
        return jsonify({"error": "invalid amount"}), 400
    if amount <= 0:
        return jsonify({"error": "amount must be positive"}), 400

    r = Retailer.query.get_or_404(retailer_id)

    created_timestamp = now_ist()
    if entry_date:
        try:
            d = date.fromisoformat(str(entry_date))
            created_timestamp = datetime.combine(d, datetime.now(IST).time())
        except ValueError:
            pass

    note = (data.get("note") or "").strip()
    t = Transaction(
        retailer_id=r.id,
        book=book,
        type=ttype,
        amount=amount,
        note=note,
        created_at=created_timestamp
    )
    db.session.add(t)
    db.session.commit()
    return jsonify({"ok": True, "balance": r.balance(book)})


# ---------------- Reports Grid & Export ----------------


def build_report_data(book, from_str=None, to_str=None):
    if not valid_book(book):
        raise ValueError("invalid book")

    today = now_ist().date()
    if from_str:
        d_from = date.fromisoformat(from_str)
    else:
        d_from = date(today.year, today.month, 1)

    if to_str:
        d_to = date.fromisoformat(to_str)
    else:
        d_to = today

    book_filter = Retailer.has_santhoor if book == "santhoor" else Retailer.has_mtr

    # Fetch all relevant retailers in 1 single query
    retailers = (
        Retailer.query.filter(book_filter == True)
        .order_by(Retailer.name.asc())
        .all()
    )
    if not retailers:
        return {
            "book": book,
            "from": d_from.isoformat(),
            "to": d_to.isoformat(),
            "rows": [],
            "totals": {"grand_credit": 0.0, "grand_debit": 0.0, "grand_balance": 0.0},
        }

    retailer_ids = [r.id for r in retailers]

    # Fetch ALL transactions up to d_to for this book in 1 single query
    dt_to_end = datetime.combine(d_to, datetime.max.time())
    all_txs = (
        Transaction.query.filter(
            Transaction.book == book,
            Transaction.retailer_id.in_(retailer_ids),
            Transaction.created_at <= dt_to_end,
        )
        .order_by(Transaction.created_at.asc())
        .all()
    )

    # Group transactions by retailer_id in memory
    txs_by_retailer = {r_id: [] for r_id in retailer_ids}
    for t in all_txs:
        txs_by_retailer[t.retailer_id].append(t)

    rows = []
    grand_credit = 0.0
    grand_debit = 0.0
    grand_balance = 0.0

    for r in retailers:
        r_txs = txs_by_retailer.get(r.id, [])

        # Balance up to d_to
        rem_balance = round(
            sum(t.amount if t.type == "purchase" else -t.amount for t in r_txs), 2
        )

        # Transactions in date range [d_from, d_to]
        range_txs = [t for t in r_txs if t.created_at.date() >= d_from]
        payments = [t for t in range_txs if t.type == "payment"]
        purchases = [t for t in range_txs if t.type == "purchase"]

        credit_total = round(sum(t.amount for t in payments), 2)
        last_credit = max((t.created_at for t in payments), default=None)
        credit_date_str = last_credit.strftime("%d %b") if last_credit else "-"

        debit_total = round(sum(t.amount for t in purchases), 2)
        last_debit = max((t.created_at for t in purchases), default=None)
        debit_date_str = last_debit.strftime("%d %b") if last_debit else "-"

        grand_credit = round(grand_credit + credit_total, 2)
        grand_debit = round(grand_debit + debit_total, 2)
        grand_balance = round(grand_balance + rem_balance, 2)

        rows.append({
            "id": r.id,
            "name": r.name,
            "credit_date": credit_date_str,
            "credit_amount": credit_total,
            "debit_date": debit_date_str,
            "debit_amount": debit_total,
            "balance": rem_balance
        })

    totals = {
        "grand_credit": grand_credit,
        "grand_debit": grand_debit,
        "grand_balance": grand_balance
    }

    return {
        "book": book,
        "from": d_from.isoformat(),
        "to": d_to.isoformat(),
        "rows": rows,
        "totals": totals
    }



@app.route("/api/reports/grid")
@app.route("/api/grid")
@login_required
def grid_report():
    book = request.args.get("book", "santhoor")
    from_str = request.args.get("from")
    to_str = request.args.get("to")
    try:
        data = build_report_data(book, from_str, to_str)
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@app.route("/api/reports/excel")
@login_required
def export_excel_report():
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"error": "openpyxl module is missing. Please run 'pip install openpyxl' or 'pip install -r requirements.txt'"}), 500

    book = request.args.get("book", "santhoor")
    from_str = request.args.get("from")
    to_str = request.args.get("to")

    try:
        report = build_report_data(book, from_str, to_str)
    except Exception as e:
        return jsonify({"error": str(e)}), 400

    wb = openpyxl.Workbook()
    ws = wb.active

    book_title = "SANTOOR PRODUCT LINE" if book == "santhoor" else "MTR PRODUCT LINE"
    ws.title = f"{book_title[:10]} Report"

    # Styling definitions
    title_font = Font(name="Arial", size=16, bold=True, color="1C1B18")
    subtitle_font = Font(name="Arial", size=11, italic=True, color="6B6558")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    bold_data_font = Font(name="Arial", size=10, bold=True)
    
    header_fill = PatternFill(start_color="0F6E6A" if book == "santhoor" else "B5680B", end_color="0F6E6A" if book == "santhoor" else "B5680B", fill_type="solid")
    total_fill = PatternFill(start_color="F2EFE9", end_color="F2EFE9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="DED7C5"),
        right=Side(style="thin", color="DED7C5"),
        top=Side(style="thin", color="DED7C5"),
        bottom=Side(style="thin", color="DED7C5")
    )
    thick_top_double_bottom = Border(
        top=Side(style="thin", color="1C1B18"),
        bottom=Side(style="double", color="1C1B18")
    )

    # Title rows
    ws.cell(row=1, column=1, value=f"CREDIT & DEBIT LEDGER REPORT - {book_title}").font = title_font
    ws.cell(row=2, column=1, value=f"Date Range: {report['from']} to {report['to']} | Generated: {now_ist().strftime('%d %b %Y %I:%M %p')} IST").font = subtitle_font

    ws.append([])  # Blank row

    # Headers
    headers = [
        "Retailer Name",
        "Credit Date (Payments)",
        "Credit Amount (₹)",
        "Debit Date (Purchases)",
        "Debit Amount (₹)",
        "Net Balance (₹)"
    ]
    ws.append(headers)

    header_row_idx = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [2, 4] else ("right" if col_idx in [3, 5, 6] else "left"), vertical="center")

    # Data Rows
    current_row = 5
    num_fmt = "#,##0.00"

    for r in report["rows"]:
        row_vals = [
            r["name"],
            r["credit_date"],
            r["credit_amount"],
            r["debit_date"],
            r["debit_amount"],
            r["balance"]
        ]

        ws.append(row_vals)
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left")
            elif col_idx in [2, 4]:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = num_fmt

        current_row += 1

    # Totals Row
    totals_vals = [
        "TOTAL",
        "-",
        report["totals"]["grand_credit"],
        "-",
        report["totals"]["grand_debit"],
        report["totals"]["grand_balance"]
    ]

    ws.append(totals_vals)
    for col_idx in range(1, len(totals_vals) + 1):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.font = bold_data_font
        cell.fill = total_fill
        cell.border = thick_top_double_bottom
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="left")
        elif col_idx in [2, 4]:
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.alignment = Alignment(horizontal="right")
            if isinstance(cell.value, (int, float)):
                cell.number_format = num_fmt

    # Adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Ledger_Report_{book}_{report['from']}_to_{report['to']}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


# ---------------- Frontend ----------------


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))

